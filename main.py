import os
import re
import requests
from bs4 import BeautifulSoup
import openpyxl
from datetime import datetime, date, time, timedelta
from ics import Calendar, Event

# ==================== KONFIGURACJA ====================
URL_STRONY_PLANU = "https://ans-elblag.pl/iis-plany-zajec.html"
INDEX_NUMBER = "21459"                                        # Podaj swój numer indeksu
TARGET_KIERUNEK = "III rok INFORMATYKA IOSI"                  # Dokładna nazwa ze strony
OUTPUT_ICS = "plan_zajec_IOSI.ics"                            # Wyjściowy plik kalendarza
CACHE_FILE = "last_update_IOSI.txt"                           # Zapis ostatniej daty modyfikacji

# Data początkowa poniedziałku pierwszego tygodnia semestru
SEMESTER_START_MONDAY = date(2026, 10, 5) 
# ======================================================

def check_and_get_excel_url(page_url, target_name):
    headers = {'User-Agent': 'Mozilla/5.0'}
    response = requests.get(page_url, headers=headers)
    soup = BeautifulSoup(response.text, 'html.parser')

    excel_url = None
    update_date = None

    for element in soup.find_all(['p', 'li', 'div']):
        text = element.get_text()
        if target_name in text:
            date_match = re.search(r'aktualizacja\s*([\d\.]+)', text, re.IGNORECASE)
            if date_match:
                update_date = date_match.group(1)

            a_tag = element.find('a', href=True)
            if a_tag:
                excel_url = a_tag['href']
                if not excel_url.startswith('http'):
                    base_url = "/".join(page_url.split('/')[:3])
                    excel_url = f"{base_url}{excel_url}"
            break

    if not excel_url:
        raise ValueError(f"Nie znaleziono pliku dla: {target_name}")

    is_updated = True
    if os.path.exists(CACHE_FILE):
        with open(CACHE_FILE, "r", encoding="utf-8") as f:
            last_date = f.read().strip()
        if last_date == update_date and update_date is not None:
            is_updated = False

    if is_updated and update_date:
        with open(CACHE_FILE, "w", encoding="utf-8") as f:
            f.write(update_date)

    return is_updated, excel_url, update_date

def get_student_groups(wb, student_index):
    groups = []
    str_idx = str(student_index).strip()

    if 'Podział na grupy' in wb.sheetnames:
        ws = wb['Podział na grupy']
        for col in range(1, ws.max_column + 1):
            group_name = None
            for row in range(1, ws.max_row + 1):
                val = str(ws.cell(row=row, column=col).value or '').strip()
                if 'gr.' in val.lower() or 'grupa' in val.lower():
                    group_name = val
                if val == str_idx and group_name:
                    groups.append(group_name)

    if 'PODZIAŁ' in wb.sheetnames:
        ws = wb['PODZIAŁ']
        for col in range(1, ws.max_column + 1):
            for row in range(1, ws.max_row + 1):
                val = str(ws.cell(row=row, column=col).value or '').strip()
                if str_idx in val:
                    header = str(ws.cell(row=3, column=col).value or '').strip()
                    if header:
                        groups.append(header)

    return list(set(groups))

def parse_schedule_directly(excel_path, my_groups):
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    plan_sheet = 'Plan' if 'Plan' in wb.sheetnames else 'plan '
    ws = wb[plan_sheet]

    day_columns = {col: i for i, col in enumerate([2, 17, 41, 56, 73])}
    time_pattern = re.compile(r'godz\.?\s*(\d{1,2}:\d{2})\s*[-–—]\s*(\d{1,2}:\d{2})', re.IGNORECASE)

    events = []

    for rng in ws.merged_cells.ranges:
        text = str(ws.cell(row=rng.min_row, column=rng.min_col).value or '').strip()
        if not text:
            continue

        match = time_pattern.search(text)
        if not match:
            continue

        start_str, end_str = match.group(1), match.group(2)
        start_h, start_m = map(int, start_str.split(':'))
        end_h, end_m = map(int, end_str.split(':'))

        day_idx = None
        for col_start in sorted(day_columns.keys()):
            if rng.min_col >= col_start:
                day_idx = day_columns[col_start]

        if day_idx is None:
            continue

        lines = [line.strip() for line in text.split('\n') if line.strip() and not time_pattern.search(line)]
        
        title = lines[0] if lines else "Zajęcia"
        lecturer = lines[1] if len(lines) > 1 else ""
        room = lines[2] if len(lines) > 2 else ""

        event_date = SEMESTER_START_MONDAY + timedelta(days=day_idx)
        start_dt = datetime.combine(event_date, time(hour=start_h, minute=start_m))
        end_dt = datetime.combine(event_date, time(hour=end_h, minute=end_m))

        events.append({
            'summary': title,
            'location': room,
            'description': f"Prowadzący: {lecturer}\nGrupy: {', '.join(my_groups)}",
            'start': start_dt,
            'end': end_dt
        })

    return events

def generate_ics(events_list, output_path):
    cal = Calendar()
    for item in events_list:
        event = Event()
        event.name = item['summary']
        event.location = item['location']
        event.description = item['description']
        event.begin = item['start']
        event.end = item['end']
        cal.events.add(event)

    with open(output_path, 'w', encoding='utf-8') as f:
        f.writelines(cal.serialize_iter())

def main():
    print(f"Sprawdzanie aktualizacji planu na stronie uczelni dla: {TARGET_KIERUNEK}...")
    try:
        updated, excel_url, update_date = check_and_get_excel_url(URL_STRONY_PLANU, TARGET_KIERUNEK)
    except Exception as e:
        print(f"Błąd sprawdzania strony: {e}")
        # Jeśli plik lokalny istnieje, sparafrazujmy lokalnie
        excel_url = None
        updated = False

    excel_file = "pobrany_plan.xlsx"

    if updated and excel_url:
        print(f"Pobieranie nowego pliku Excel z dnia {update_date}...")
        res = requests.get(excel_url)
        with open(excel_file, 'wb') as f:
            f.write(res.content)
    elif not os.path.exists(excel_file):
        print("Brak zaktualizowanej daty, ale plik lokalny nie istnieje. Pobieganie wymuszone...")
        if excel_url:
            res = requests.get(excel_url)
            with open(excel_file, 'wb') as f:
                f.write(res.content)

    if os.path.exists(excel_file):
        wb = openpyxl.load_workbook(excel_file, data_only=True)
        my_groups = get_student_groups(wb, INDEX_NUMBER)
        print(f"Znalezione grupy dla Twojego indeksu ({INDEX_NUMBER}): {my_groups}")

        events = parse_schedule_directly(excel_file, my_groups)
        print(f"Przetworzono {len(events)} zajęć.")

        generate_ics(events, OUTPUT_ICS)
        print(f"Wygenerowano plik {OUTPUT_ICS}!")

if __name__ == "__main__":
    main()
